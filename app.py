import io
import os
import re
import tempfile
import traceback
import struct
import zipfile
import zlib
from urllib.parse import urljoin

import pandas as pd
import requests
from PIL import Image, ImageDraw, ImageOps
import streamlit as st
from bs4 import BeautifulSoup

try:
    import pdfplumber
except Exception as e:
    pdfplumber = None
    PDF_IMPORT_ERROR = str(e)

try:
    from pdf2docx import Converter
except Exception as e:
    Converter = None
    PDF2DOCX_ERROR = str(e)

try:
    from docx import Document
except Exception as e:
    Document = None
    DOCX_ERROR = str(e)

try:
    import pymupdf as fitz
except Exception:
    try:
        import fitz
    except Exception:
        fitz = None

try:
    import pypdfium2 as pdfium
except Exception as e:
    pdfium = None
    PDFIUM_ERROR = str(e)

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                  "(KHTML, like Gecko) Chrome/124.0 Safari/537.36",
    "Accept": "text/html,application/pdf,*/*",
}
DEFAULT_URL = "https://krg.ac.in/notice/355/"


def fetch(url, timeout=60):
    r = requests.get(url, headers=HEADERS, timeout=timeout, verify=False)
    r.raise_for_status()
    return r


def get_notice_info(url):
    html = fetch(url).text
    soup = BeautifulSoup(html, "html.parser")

    title = ""
    for h1 in soup.find_all("h1"):
        t = h1.get_text(strip=True)
        if t and t.lower() != "notice viewer" and "college" not in t.lower():
            title = t
            break

    m = re.search(r"Published on\s+([^\n<]+)", soup.get_text("\n"))
    date = m.group(1).strip() if m else ""

    pdf_url = ""
    found = re.findall(r"""["'(=\s]([^"'\s()<>]*?/media/notices/[^"'\s()<>]*?\.pdf)""", html, flags=re.I)
    if found:
        pdf_url = urljoin(url, found[0])
    else:
        for tag, attr in (("a", "href"), ("iframe", "src"), ("embed", "src"), ("object", "data")):
            for el in soup.find_all(tag):
                link = el.get(attr, "")
                if ".pdf" in link.lower() and "/static/images/" not in link.lower():
                    pdf_url = urljoin(url, link)
                    break
            if pdf_url:
                break
    return {"title": title, "date": date, "page_url": url, "pdf_url": pdf_url}


def _norm(x):
    return re.sub(r"[^a-z0-9\u0900-\u097f]", "", str(x).lower())


def std_header(h):
    """Alag-alag PDF ke header ko ek jaise naam do, taki combine karne par data ek hi column me aaye."""
    n = _norm(h)
    if not n:
        return ""
    if n in ("sno", "srno", "slno", "serialno", "serialnumber", "sr", "sl", "no", "क्र", "क्रमांक", "क्रसं", "क्रम") \
            or n.startswith(("sno", "srno", "slno", "serialno", "क्र")):
        return "S. No."
    if "father" in n or "mother" in n or "guardian" in n or "पिता" in n or "माता" in n:
        return "Father/Mother Name"
    if "name" in n or "नाम" in n:
        return "Student Name"
    if "enrol" in n or "enrl" in n:
        return "Enrollment No."
    if "roll" in n or "रोल" in n:
        return "Roll No."
    return str(h).strip()


def _is_header_row(row):
    hit = {std_header(c) for c in row}
    return "Student Name" in hit or ("S. No." in hit and len(hit & {"Roll No.", "Enrollment No.", "Father/Mother Name"}) > 0)


def extract_table(pdf_bytes):
    if pdfplumber is None:
        raise RuntimeError(f"pdfplumber import nahi hua: {PDF_IMPORT_ERROR}")
    rows = []
    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        for page in pdf.pages:
            tables = page.extract_tables()
            if tables:
                for table in tables:
                    for row in table:
                        rows.append([(c or "").replace("\n", " ").strip() for c in row])
            else:
                for line in (page.extract_text() or "").splitlines():
                    parts = re.split(r"\s{2,}", line.strip())
                    if parts and any(parts):
                        rows.append(parts)

    rows = [r for r in rows if any(c for c in r)]
    if not rows:
        raise RuntimeError("PDF me text/table nahi mili (scanned image PDF ho sakti hai).")

    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]

    # Asli header row dhundo (title/college ke naam wali upar ki rows ko chhodo)
    h_idx = next((i for i, r in enumerate(rows[:25]) if _is_header_row(r)), None)
    if h_idx is not None:
        header = [std_header(c) for c in rows[h_idx]]
        body = rows[h_idx + 1:]
    elif rows[0][0].strip().isdigit():
        # Header hai hi nahi, seedha data: pehla column S. No., doosra Student Name maano
        header = ["S. No.", "Student Name"] + [f"Column {i + 1}" for i in range(2, width)]
        body = rows
    else:
        header = [std_header(c) for c in rows[0]]
        body = rows[1:]

    header = [h if h else f"Column {i + 1}" for i, h in enumerate(header)]
    seen, uniq = {}, []
    for h in header:
        seen[h] = seen.get(h, 0) + 1
        uniq.append(h if seen[h] == 1 else f"{h}_{seen[h]}")

    # Har page par dobara aane wali header rows hatao
    hn = [_norm(std_header(c)) for c in header]
    body = [r for r in body if [_norm(std_header(c)) for c in r] != hn]
    return pd.DataFrame(body, columns=uniq)


def add_title(df, title):
    df = df.copy()
    df["Title"] = title
    return df


def to_excel_bytes(df, status_df):
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="All Data", index=False)
        status_df.to_excel(writer, sheet_name="Notice Info", index=False)
        for ws in writer.book.worksheets:
            for col in ws.columns:
                longest = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[col[0].column_letter].width = min(longest + 2, 60)
    return buf.getvalue()


def show_result(frames, status_rows):
    status_df = pd.DataFrame(status_rows)
    st.subheader("Status")
    st.dataframe(status_df, use_container_width=True)
    if not frames:
        st.error("Kisi bhi link/PDF se data nahi mila.")
        return
    big = pd.concat(frames, ignore_index=True).fillna("")
    st.success(f"Total {len(big)} rows, {len(frames)} notice se")
    st.dataframe(big, use_container_width=True)
    st.download_button(
        "⬇️ Excel download karo",
        data=to_excel_bytes(big, status_df),
        file_name="notices_data.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def safe_name(text, default="notice"):
    return re.sub(r"[^A-Za-z0-9._-]+", "_", text or default).strip("_") or default


def pdf_to_docx_bytes(pdf_bytes):
    """Original PDF jaisa layout rakhte hue PDF ko Word (.docx) me badlo."""
    if Converter is None:
        raise RuntimeError(f"pdf2docx install nahi hai: {PDF2DOCX_ERROR}")
    with tempfile.TemporaryDirectory() as tmp:
        pdf_path = os.path.join(tmp, "in.pdf")
        docx_path = os.path.join(tmp, "out.docx")
        with open(pdf_path, "wb") as f:
            f.write(pdf_bytes)
        cv = Converter(pdf_path)
        try:
            cv.convert(docx_path)
        finally:
            cv.close()
        with open(docx_path, "rb") as f:
            return f.read()


def tables_to_docx_bytes(frames_with_titles):
    """Har notice ki table ko Word table me daalo (title heading ke saath), sab ek hi file me."""
    if Document is None:
        raise RuntimeError(f"python-docx install nahi hai: {DOCX_ERROR}")
    doc = Document()
    for n, (title, df) in enumerate(frames_with_titles):
        if n:
            doc.add_page_break()
        doc.add_heading(title or "Notice", level=1)
        cols = [c for c in df.columns if c != "Title"]
        table = doc.add_table(rows=1, cols=len(cols))
        table.style = "Table Grid"
        for i, c in enumerate(cols):
            table.rows[0].cells[i].text = str(c)
            for p in table.rows[0].cells[i].paragraphs:
                for r in p.runs:
                    r.bold = True
        for _, row in df.iterrows():
            cells = table.add_row().cells
            for i, c in enumerate(cols):
                cells[i].text = str(row[c])
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_zip(files):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in files:
            z.writestr(name, data)
    return buf.getvalue()


def download_pdfs_from_links(text):
    """Links se PDF download karo. Return: [(title, pdf_bytes)], status_rows"""
    urls = [u.strip() for u in text.splitlines() if u.strip()]
    pdfs, status = [], []
    bar = st.progress(0.0)
    for i, u in enumerate(urls, 1):
        row = {"Link": u, "Title": "", "Result": ""}
        try:
            info = get_notice_info(u)
            row["Title"] = info["title"]
            if not info["pdf_url"]:
                raise RuntimeError("PDF link nahi mila")
            pdfs.append((info["title"] or f"notice_{i}", fetch(info["pdf_url"]).content))
            row["Result"] = "PDF mil gayi"
        except Exception as e:
            row["Result"] = f"Error: {e}"
        status.append(row)
        bar.progress(i / len(urls))
    return pdfs, status


def convert_and_offer(pdfs, status, zip_name):
    """PDFs ko Word me badlo aur download button do (ek ho to .docx, zyada ho to .zip)."""
    outputs = []
    for title, data in pdfs:
        row = next((r for r in status if r.get("Title") == title or r.get("Link") == title), None)
        try:
            outputs.append((safe_name(title) + ".docx", pdf_to_docx_bytes(data)))
            if row is not None:
                row["Result"] = "Word ban gaya"
        except Exception as e:
            if row is not None:
                row["Result"] = f"Error: {e}"
            else:
                status.append({"Link": title, "Title": title, "Result": f"Error: {e}"})
    st.subheader("Status")
    st.dataframe(pd.DataFrame(status), use_container_width=True)
    DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if not outputs:
        st.error("Koi Word file nahi bani.")
    elif len(outputs) == 1:
        st.download_button("⬇️ Word download karo", data=outputs[0][1], file_name=outputs[0][0], mime=DOCX_MIME)
    else:
        st.download_button("⬇️ Sab Word files (ZIP) download karo", data=make_zip(outputs),
                           file_name=zip_name, mime="application/zip")



# ======================= SIZE BADLO / FORMAT BADLO =======================
UNITS = {"KB": 1024, "MB": 1024 ** 2, "GB": 1024 ** 3}
MAX_PAD = 200 * 1024 ** 2  # size badhane ki seema (server memory ke liye)
MIME = {
    "pdf": "application/pdf", "jpg": "image/jpeg", "jpeg": "image/jpeg", "png": "image/png",
    "webp": "image/webp", "bmp": "image/bmp", "tiff": "image/tiff", "zip": "application/zip",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def human(n):
    n = float(n)
    for u in ("B", "KB", "MB", "GB"):
        if n < 1024 or u == "GB":
            return f"{n:.0f} {u}" if u == "B" else f"{n:.2f} {u}"
        n /= 1024


def _open_img(data):
    im = Image.open(io.BytesIO(data))
    im.load()
    return im


def _flatten(im):
    """Transparent image ko safed background par RGB banao (JPG/BMP ke liye)."""
    if im.mode in ("RGBA", "LA") or (im.mode == "P" and "transparency" in im.info):
        im = im.convert("RGBA")
        bg = Image.new("RGB", im.size, (255, 255, 255))
        bg.paste(im, mask=im.split()[3])
        return bg
    return im.convert("RGB")


def _jpeg_bytes(im, q):
    b = io.BytesIO()
    im.save(b, "JPEG", quality=q, optimize=True)
    return b.getvalue()


def _png_bytes(im, quantize=False):
    b = io.BytesIO()
    if quantize:
        if im.mode in ("RGBA", "LA", "P"):
            im = im.convert("RGBA").quantize(colors=256, method=Image.Quantize.FASTOCTREE)
        else:
            im = im.convert("RGB").quantize(colors=256)
    im.save(b, "PNG", optimize=True)
    return b.getvalue()


def _scaled(im, scale):
    if scale == 1.0:
        return im
    return im.resize((max(1, int(im.width * scale)), max(1, int(im.height * scale))), Image.LANCZOS)


def shrink_image(data, fmt, target):
    im = _open_img(data)
    scale = 1.0
    if fmt == "JPEG":
        base = _flatten(im)
        while True:
            cur = _scaled(base, scale)
            lo, hi, best = 5, 95, None
            while lo <= hi:
                mid = (lo + hi) // 2
                out = _jpeg_bytes(cur, mid)
                if len(out) <= target:
                    best, lo = out, mid + 1
                else:
                    hi = mid - 1
            if best:
                return best
            scale *= 0.85
            if min(base.width, base.height) * scale < 24:
                return _jpeg_bytes(cur, 5)
    while True:  # PNG
        cur = _scaled(im, scale)
        out = _png_bytes(cur, False)
        if len(out) <= target:
            return out
        out = _png_bytes(cur, True)
        if len(out) <= target:
            return out
        scale *= 0.9
        if min(im.width, im.height) * scale < 24:
            return out


def pad_jpeg(data, extra):
    pos = 2
    if data[2:4] == b"\xff\xe0":  # JFIF APP0 ke baad comment jodo
        pos = 4 + struct.unpack(">H", data[4:6])[0]
    segs, remaining = bytearray(), extra
    while remaining >= 5:
        p = min(65533, remaining - 4)
        segs += b"\xff\xfe" + struct.pack(">H", p + 2) + bytes(p)
        remaining -= p + 4
    return data[:pos] + bytes(segs) + data[pos:] + bytes(remaining)


def pad_png(data, extra):
    if extra < 20 or data[-12:-8] != b"\x00\x00\x00\x00" or data[-8:-4] != b"IEND":
        return data + bytes(extra)
    payload = b"Padding\x00" + bytes(extra - 12 - 8)
    chunk = struct.pack(">I", len(payload)) + b"tEXt" + payload
    chunk += struct.pack(">I", zlib.crc32(b"tEXt" + payload) & 0xFFFFFFFF)
    return data[:-12] + chunk + data[-12:]


def pad_pdf(data, extra):
    if fitz is not None and extra > 4000:
        try:
            doc = fitz.open(stream=data, filetype="pdf")
            doc.embfile_add("padding.bin", os.urandom(extra - 3000))
            out = doc.tobytes()
            doc.close()
            if len(out) <= len(data) + extra:
                return out + b"\n%" + b"0" * (len(data) + extra - len(out) - 2) if len(data) + extra - len(out) > 2 else out
        except Exception:
            pass
    return data + b"\n%" + b"0" * (extra - 2) if extra > 2 else data + bytes(extra)


def _zip_rebuild(data, media_fn=None):
    zin = zipfile.ZipFile(io.BytesIO(data))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as zout:
        for item in zin.infolist():
            b = zin.read(item.filename)
            low = item.filename.lower()
            if media_fn and "/media/" in low and low.endswith((".jpg", ".jpeg", ".png")):
                try:
                    b = media_fn(low, b)
                except Exception:
                    pass
            zout.writestr(item.filename, b)
    return buf.getvalue()


def _recompress_media(name, b, scale, q):
    im = _open_img(b)
    cur = _scaled(im, scale)
    if name.endswith(".png"):
        out = _png_bytes(cur, quantize=(scale < 1 or q < 60))
    else:
        out = _jpeg_bytes(_flatten(cur), q)
    return out if len(out) < len(b) else b


def shrink_office(data, target):
    best = _zip_rebuild(data)
    if len(best) <= target:
        return best, "Lossless compress"
    for scale, q in [(1, 85), (1, 70), (.8, 60), (.7, 50), (.6, 40), (.5, 30), (.4, 25), (.3, 20)]:
        out = _zip_rebuild(data, lambda n, b: _recompress_media(n, b, scale, q))
        if len(out) < len(best):
            best = out
        if len(out) <= target:
            return out, f"Andar ki images compress ki (quality {q}, size {int(scale * 100)}%)"
    return best, "images compress karne ke baad bhi target tak nahi pahunch paya"


def pad_office(data, extra):
    zin = zipfile.ZipFile(io.BytesIO(data))

    def build(pad_len):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as zout:
            for item in zin.infolist():
                b = zin.read(item.filename)
                if item.filename == "[Content_Types].xml" and b'Extension="bin"' not in b and b"<Default " in b:
                    b = b.replace(b"<Default ", b'<Default Extension="bin" ContentType="application/octet-stream"/><Default ', 1)
                zout.writestr(item.filename, b)
            if pad_len is not None:
                zi = zipfile.ZipInfo("customXml/padding.bin")
                zi.compress_type = zipfile.ZIP_STORED
                zout.writestr(zi, bytes(pad_len))
        return buf.getvalue()

    target = len(data) + extra
    s0, s1 = len(build(None)), len(build(1))
    header = s1 - s0 - 1
    pad_len = max(target - s0 - header, 0)
    return build(pad_len)


def shrink_pdf(data, target):
    best, note = data, ""
    if fitz is not None:
        try:
            doc = fitz.open(stream=data, filetype="pdf")
            out = doc.tobytes(garbage=4, deflate=True, clean=True)
            doc.close()
            if len(out) <= target:
                return out, "Lossless compress (text jaisa ka taisa)"
            if len(out) < len(best):
                best = out
        except Exception:
            pass
    if pdfium is None:
        return best, f"pypdfium2 nahi hai: {PDFIUM_ERROR}"
    pdf = pdfium.PdfDocument(data)
    base = [pdf[i].render(scale=150 / 72).to_pil().convert("RGB") for i in range(len(pdf))]
    for dpi, q in [(150, 75), (120, 60), (100, 50), (80, 40), (60, 30), (50, 25)]:
        imgs = [_scaled(im, dpi / 150) for im in base]
        buf = io.BytesIO()
        imgs[0].save(buf, "PDF", save_all=True, append_images=imgs[1:], quality=q, resolution=dpi)
        out = buf.getvalue()
        if len(out) < len(best):
            best = out
        if len(out) <= target:
            return out, f"Pages image me badal kar compress kiya ({dpi} dpi, quality {q}). Ab text select/copy nahi hoga."
    return best, "bahut compress karne ke baad bhi target tak nahi pahunch paya"


def resize_file(name, data, target):
    """File ko target bytes ke paas laao. Return: (new_bytes, note)"""
    ext = name.rsplit(".", 1)[-1].lower()
    cur = len(data)
    if abs(target - cur) <= 16:
        return data, "Size pehle se hi itna hai"
    if target > cur:
        extra = target - cur
        if extra > MAX_PAD:
            raise RuntimeError(f"Size badhane ki seema {human(MAX_PAD)} hai")
        if ext in ("jpg", "jpeg"):
            return pad_jpeg(data, extra), "Size badhaya (extra bytes jode, quality same)"
        if ext == "png":
            return pad_png(data, extra), "Size badhaya (extra bytes jode, quality same)"
        if ext == "pdf":
            return pad_pdf(data, extra), "Size badhaya (extra bytes jode, content same)"
        if ext in ("docx", "xlsx"):
            return pad_office(data, extra), "Size badhaya (extra data jodkar, content same)"
        raise RuntimeError("Ye file type support nahi hai")
    if ext in ("jpg", "jpeg"):
        return shrink_image(data, "JPEG", target), "Compress kiya"
    if ext == "png":
        return shrink_image(data, "PNG", target), "Compress kiya"
    if ext == "pdf":
        return shrink_pdf(data, target)
    if ext in ("docx", "xlsx"):
        return shrink_office(data, target)
    raise RuntimeError("Ye file type support nahi hai")


IMG_FORMATS = {"JPG": ("JPEG", "jpg"), "PNG": ("PNG", "png"), "WEBP": ("WEBP", "webp"),
               "BMP": ("BMP", "bmp"), "TIFF": ("TIFF", "tiff")}


def convert_image(data, key):
    fmt, _ = IMG_FORMATS[key]
    im = _open_img(data)
    if fmt in ("JPEG", "BMP"):
        im = _flatten(im)
    elif im.mode not in ("RGB", "RGBA", "L", "P"):
        im = im.convert("RGBA")
    buf = io.BytesIO()
    kw = {"quality": 95} if fmt in ("JPEG", "WEBP") else {}
    if fmt in ("JPEG", "PNG"):
        kw["optimize"] = True
    im.save(buf, fmt, **kw)
    return buf.getvalue()


def images_to_pdf(images_bytes):
    imgs = [_flatten(_open_img(b)) for b in images_bytes]
    buf = io.BytesIO()
    imgs[0].save(buf, "PDF", save_all=True, append_images=imgs[1:], resolution=100)
    return buf.getvalue()


def pdf_to_images(pdf_bytes, key, dpi):
    if pdfium is None:
        raise RuntimeError(f"pypdfium2 install nahi hai: {PDFIUM_ERROR}")
    fmt, ext = IMG_FORMATS[key]
    pdf = pdfium.PdfDocument(pdf_bytes)
    out = []
    for i in range(len(pdf)):
        im = pdf[i].render(scale=dpi / 72).to_pil()
        im = _flatten(im) if fmt == "JPEG" else im.convert("RGB")
        buf = io.BytesIO()
        im.save(buf, fmt, **({"quality": 92} if fmt == "JPEG" else {}))
        out.append((i + 1, ext, buf.getvalue()))
    return out


def offer_download(files, zip_name, label="Download karo"):
    """files = [(name, bytes)]. Ek ho to seedha file, zyada ho to ZIP."""
    if not files:
        return
    if len(files) == 1:
        name, data = files[0]
        st.download_button(label, data=data, file_name=name,
                           mime=MIME.get(name.rsplit(".", 1)[-1].lower(), "application/octet-stream"),
                           key=f"dl_{zip_name}_{name}")
    else:
        st.download_button(f"{label} (ZIP, {len(files)} files)", data=make_zip(files), file_name=zip_name,
                           mime="application/zip", key=f"dl_{zip_name}")



# ============================ PASSPORT PHOTO ============================
PHOTO_SIZES = {
    "35 × 45 mm (sabse common)": (35, 45),
    "2 × 2 inch / 51 × 51 mm": (50.8, 50.8),
    "PAN card 25 × 35 mm": (25, 35),
    "Stamp size 20 × 25 mm": (20, 25),
    "33 × 48 mm": (33, 48),
    "Apna size (mm me)": None,
}
PHOTO_BG = {"Safed": "#FFFFFF", "Neela": "#2F6FB5", "Halka grey": "#E6E8EB", "Apna rang": None}
PAPERS = {"4 × 6 inch": (101.6, 152.4), "5 × 7 inch": (127.0, 177.8), "A4": (210.0, 297.0)}


def _flatten_on(im, bg):
    if im.mode in ("RGBA", "LA") or (im.mode == "P" and "transparency" in im.info):
        im = im.convert("RGBA")
        canvas = Image.new("RGB", im.size, bg)
        canvas.paste(im, mask=im.split()[3])
        return canvas
    return im.convert("RGB")


def make_passport(img, w_mm, h_mm, zoom, off_x, off_y, bg, dpi=300):
    """Photo ko w_mm x h_mm me katkar do. zoom<1 par kinare bg rang se bhar jate hain."""
    W, H = round(w_mm / 25.4 * dpi), round(h_mm / 25.4 * dpi)
    src = _flatten_on(img, bg)
    aspect = w_mm / h_mm
    bh = min(src.height, src.width / aspect) / zoom
    bw = bh * aspect
    cx = src.width / 2 + off_x / 100 * (src.width / 2)
    cy = src.height / 2 + off_y / 100 * (src.height / 2)
    x0, y0 = int(round(cx - bw / 2)), int(round(cy - bh / 2))
    canvas = Image.new("RGB", (max(1, int(round(bw))), max(1, int(round(bh)))), bg)
    canvas.paste(src, (-x0, -y0))
    return canvas.resize((W, H), Image.LANCZOS)


def make_sheet(photo, paper, copies=None, guides=True, dpi=300, margin_mm=2.0, gap_mm=1.0):
    """Photo ki copies ek print sheet par lagao. Return: (sheet_image, max_fit)"""
    pw_mm, ph_mm = PAPERS[paper]
    mm = lambda v: round(v / 25.4 * dpi)
    w, h = photo.size

    def fit(PW, PH):
        cols = int((mm(PW) - 2 * mm(margin_mm) + mm(gap_mm)) // (w + mm(gap_mm)))
        rows = int((mm(PH) - 2 * mm(margin_mm) + mm(gap_mm)) // (h + mm(gap_mm)))
        return max(cols, 0), max(rows, 0)

    c1, r1 = fit(pw_mm, ph_mm)
    c2, r2 = fit(ph_mm, pw_mm)
    if c2 * r2 > c1 * r1:
        PW, PH, cols, rows = ph_mm, pw_mm, c2, r2
    else:
        PW, PH, cols, rows = pw_mm, ph_mm, c1, r1
    if cols * rows == 0:
        raise RuntimeError("Is paper par ek bhi photo nahi aati. Bada paper chuniye.")
    n = cols * rows if copies is None else max(1, min(int(copies), cols * rows))
    gap = mm(gap_mm)
    used_cols, used_rows = min(cols, n), -(-n // cols)
    grid_w = used_cols * w + (used_cols - 1) * gap
    grid_h = used_rows * h + (used_rows - 1) * gap
    sheet = Image.new("RGB", (mm(PW), mm(PH)), "white")
    ox, oy = (sheet.width - grid_w) // 2, (sheet.height - grid_h) // 2
    draw = ImageDraw.Draw(sheet)
    for i in range(n):
        r, c = divmod(i, cols)
        x, y = ox + c * (w + gap), oy + r * (h + gap)
        sheet.paste(photo, (x, y))
        if guides:
            draw.rectangle([x - 1, y - 1, x + w, y + h], outline="#9AA3B2")
    return sheet, cols * rows


def _jpg(im, dpi=300, q=95):
    b = io.BytesIO()
    im.save(b, "JPEG", quality=q, dpi=(dpi, dpi), optimize=True)
    return b.getvalue()


def _pdf(im, dpi=300):
    b = io.BytesIO()
    im.save(b, "PDF", resolution=dpi)
    return b.getvalue()


st.set_page_config(page_title="Notice se Excel aur Word", page_icon="📄", layout="centered")

CSS = """
<style>
@import url('https://fonts.googleapis.com/css2?family=Bricolage+Grotesque:opsz,wght@12..96,600;12..96,800&family=Hind:wght@400;500;600&display=swap');

:root{
  --ink:#14213D; --muted:#5B6678; --paper:#F4F6F9; --card:#FFFFFF; --line:#D9DFE8;
  --xl:#1D6F42; --wd:#2B579A; --pdf:#B3261E; --teal:#0E7C86; --amber:#B26A00;
}
html, body, .stApp, [class*="css"]{ font-family:'Hind','Noto Sans Devanagari',sans-serif; color:var(--ink); }
.stApp{ background:var(--paper); }
.block-container{ max-width:780px; padding-top:1.8rem; padding-bottom:4rem; }
#MainMenu, footer{ visibility:hidden; }

/* file-shaped badge: poore app ka ek hi motif */
.fs{ flex:none; display:flex; align-items:flex-end; width:46px; height:56px; padding:0 0 6px 6px;
     background:var(--c); color:#fff; font:600 .7rem/1.05 'Hind',sans-serif; border-radius:3px;
     clip-path:polygon(0 0,68% 0,100% 24%,100% 100%,0 100%); }
.fs.big{ width:56px; height:68px; font-size:.82rem; padding:0 0 7px 8px; }

/* hero */
.hero{ margin:.2rem 0 1.8rem 0; }
.hero .files{ position:relative; height:88px; margin-bottom:.5rem; }
.hero .fs{ position:absolute; top:6px; }
.hero .fs:nth-child(1){ left:8px;  transform:rotate(-6deg); }
.hero .fs:nth-child(2){ left:58px; top:10px; transform:rotate(3deg); }
.hero .fs:nth-child(3){ left:108px; top:4px; transform:rotate(-2deg); }
.hero h1{ font-family:'Bricolage Grotesque',sans-serif; font-weight:800; font-size:2.4rem; line-height:1.07;
          letter-spacing:-.8px; margin:0 0 .65rem 0; padding:0; max-width:680px; text-wrap:balance; }
.hero p{ font-size:1.07rem; line-height:1.5; color:var(--muted); max-width:520px; margin:0; }

/* tool tiles: upar jaankari, neeche rang wala "Kholo" button */
[class*="st-key-tile_"]{ gap:0 !important; margin-bottom:.4rem; }
.tile{ display:flex; align-items:center; gap:1rem; background:var(--card); border:1px solid var(--line); border-bottom:0;
       border-radius:14px 14px 0 0; padding:1rem 1.1rem; min-height:94px; border-left:5px solid var(--c); }
.tile b{ display:block; font-family:'Bricolage Grotesque',sans-serif; font-weight:600; font-size:1.14rem; line-height:1.2; }
.tile .d{ display:block; margin-top:.2rem; color:var(--muted); font-size:.93rem; line-height:1.35; }
[class*="st-key-open_"] .stButton, [class*="st-key-open_"] .stButton > button{ width:100%; }
[class*="st-key-open_"] .stButton > button{ border:0; border-radius:0 0 14px 14px; color:#fff; font-weight:600; min-height:2.6rem; }
[class*="st-key-open_"] .stButton > button:hover{ filter:brightness(.88); color:#fff; }
.st-key-open_excel button{ background:var(--xl) !important; }
.st-key-open_word button{ background:var(--wd) !important; }
.st-key-open_pdf button{ background:var(--pdf) !important; }
.st-key-open_size button{ background:var(--ink) !important; }
.st-key-open_convert button{ background:var(--teal) !important; }
.st-key-open_photo button{ background:var(--amber) !important; }

/* tool page */
.st-key-back button{ background:none; border:0; color:var(--muted); padding:0; min-height:0; font-weight:600; }
.st-key-back button:hover{ color:var(--ink); text-decoration:underline; background:none; }
.thead{ display:flex; align-items:center; gap:1rem; margin:.7rem 0 1.1rem 0; }
.thead h2{ font-family:'Bricolage Grotesque',sans-serif; font-weight:800; font-size:1.9rem; letter-spacing:-.5px; line-height:1.1; margin:0; padding:0; }
.thead p{ margin:.25rem 0 0 0; color:var(--muted); font-size:1rem; }
.st-key-panel{ background:var(--card); border:1px solid var(--line); border-radius:14px; padding:1.3rem 1.4rem 1.5rem 1.4rem; }

/* inputs */
.stTextArea textarea{ background:#FBFCFE; border:1px solid var(--line); border-radius:10px;
     font-family:ui-monospace,Menlo,Consolas,monospace; font-size:.84rem; line-height:1.55; }
.stTextArea textarea:focus{ border-color:var(--ink); box-shadow:0 0 0 2px rgba(20,33,61,.12); }
[data-testid="stFileUploaderDropzone"]{ background:#FBFCFE; border:1.5px dashed #9AA7BD; border-radius:12px; }
div[role="radiogroup"]{ gap:.5rem; margin-bottom:.3rem; }
.hint{ color:var(--muted); font-size:.95rem; margin:.1rem 0 .9rem 0; }

/* buttons: har tool ka apna rang */
.stButton > button, .stDownloadButton > button{ border-radius:10px; font-weight:600; padding:.55rem 1.3rem; min-height:2.8rem; }
[class*="st-key-go_x"] button{ background:var(--xl);   border:1px solid var(--xl);   color:#fff; }
[class*="st-key-go_w"] button{ background:var(--wd);   border:1px solid var(--wd);   color:#fff; }
[class*="st-key-go_p"] button{ background:var(--pdf);  border:1px solid var(--pdf);  color:#fff; }
[class*="st-key-go_s"] button{ background:var(--ink);  border:1px solid var(--ink);  color:#fff; }
[class*="st-key-go_c"] button{ background:var(--teal); border:1px solid var(--teal); color:#fff; }
[class*="st-key-go_"] button:hover{ filter:brightness(.88); color:#fff; }
.stDownloadButton > button{ background:var(--card); border:1.5px solid var(--ink); color:var(--ink); }
.stDownloadButton > button:hover{ background:var(--ink); color:#fff; border-color:var(--ink); }

/* results */
[data-testid="stDataFrame"]{ border:1px solid var(--line); border-radius:10px; overflow:hidden; }
[data-testid="stAlert"]{ border-radius:10px; }
h3{ font-family:'Bricolage Grotesque',sans-serif; font-weight:600; letter-spacing:-.2px; }

@media (max-width:640px){
  .hero h1{ font-size:1.9rem; }
  .thead h2{ font-size:1.55rem; }
  .st-key-panel{ padding:1rem 1rem 1.2rem 1rem; }
}
</style>
"""
st.markdown(CSS, unsafe_allow_html=True)

SRC = ["Notice links", "PDF upload"]
LINKS_LABEL = "Notice links"
LINKS_HINT = '<div class="hint">Har line me ek link. 10 ya usse zyada bhi daal sakte hain.</div>'


def ui_links_to_excel():
    st.markdown(LINKS_HINT, unsafe_allow_html=True)
    links = st.text_area(LINKS_LABEL, value=DEFAULT_URL, height=200, key="links_x", label_visibility="collapsed")
    if st.button("Excel banao", key="go_x_links"):
        urls = [u.strip() for u in links.splitlines() if u.strip()]
        if not urls:
            st.error("Pehle kam se kam ek notice link daaliye.")
            return
        frames, status_rows = [], []
        bar = st.progress(0.0)
        for i, u in enumerate(urls, 1):
            row = {"Link": u, "Title": "", "Rows": 0, "Result": ""}
            try:
                info = get_notice_info(u)
                row["Title"] = info["title"]
                if not info["pdf_url"]:
                    raise RuntimeError("PDF link nahi mila")
                df = extract_table(fetch(info["pdf_url"]).content)
                frames.append(add_title(df, info["title"]))
                row["Rows"] = len(df)
                row["Result"] = "OK"
            except Exception as e:
                row["Result"] = f"Error: {e}"
            status_rows.append(row)
            bar.progress(i / len(urls))
        show_result(frames, status_rows)


def ui_pdf_to_excel():
    ups = st.file_uploader("PDF chuniye (ek ya zyada)", type=["pdf"], accept_multiple_files=True, key="up_x")
    if ups and st.button("Excel banao", key="go_x_pdf"):
        frames, status_rows = [], []
        for up in ups:
            title = up.name.rsplit(".", 1)[0]
            row = {"Link": up.name, "Title": title, "Rows": 0, "Result": ""}
            try:
                df = extract_table(up.read())
                frames.append(add_title(df, title))
                row["Rows"] = len(df)
                row["Result"] = "OK"
            except Exception as e:
                row["Result"] = f"Error: {e}"
            status_rows.append(row)
        show_result(frames, status_rows)


def ui_links_to_word():
    st.markdown(LINKS_HINT, unsafe_allow_html=True)
    links = st.text_area(LINKS_LABEL, value=DEFAULT_URL, height=200, key="links_w", label_visibility="collapsed")
    if st.button("Word banao", key="go_w_links"):
        if not links.strip():
            st.error("Pehle kam se kam ek notice link daaliye.")
            return
        pdfs, status = download_pdfs_from_links(links)
        with st.spinner("Word me convert ho raha hai..."):
            convert_and_offer(pdfs, status, "notices_word.zip")


def ui_pdf_to_word():
    ups = st.file_uploader("PDF chuniye (ek ya zyada)", type=["pdf"], accept_multiple_files=True, key="up_w")
    if ups and st.button("Word banao", key="go_w_pdf"):
        pdfs = [(u.name.rsplit(".", 1)[0], u.read()) for u in ups]
        status = [{"Link": t, "Title": t, "Result": ""} for t, _ in pdfs]
        with st.spinner("Word me convert ho raha hai..."):
            convert_and_offer(pdfs, status, "pdf_to_word.zip")


def ui_table_to_word():
    st.markdown(
        '<div class="hint">Saare notice ek hi Word file me aayenge. Har notice naye page par, apne title ke saath.</div>',
        unsafe_allow_html=True,
    )
    links = st.text_area(LINKS_LABEL, value="", height=130, key="links_t", label_visibility="collapsed",
                         placeholder="Notice links yahan daaliye (har line me ek)")
    ups = st.file_uploader("Ya PDF chuniye", type=["pdf"], accept_multiple_files=True, key="up_t")
    if st.button("Word banao", key="go_w_table"):
        items, status = [], []
        if links.strip():
            pdfs, status = download_pdfs_from_links(links)
            items += pdfs
        items += [(u.name.rsplit(".", 1)[0], u.read()) for u in (ups or [])]
        if not items:
            st.error("Link daaliye ya PDF chuniye.")
            return
        results = []
        for title, data in items:
            try:
                results.append((title, extract_table(data)))
            except Exception as e:
                status.append({"Link": title, "Title": title, "Result": f"Error: {e}"})
        if status:
            st.dataframe(pd.DataFrame(status), use_container_width=True)
        if results:
            st.success(f"{len(results)} notice ki table taiyar hai")
            st.download_button(
                "Word download karo",
                data=tables_to_docx_bytes(results),
                file_name="notices_tables.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        else:
            st.error("Kisi bhi notice se table nahi mili.")


def ui_links_to_pdf():
    st.markdown(LINKS_HINT, unsafe_allow_html=True)
    links = st.text_area(LINKS_LABEL, value=DEFAULT_URL, height=200, key="links_p", label_visibility="collapsed")
    if st.button("PDF nikalo", key="go_p_links"):
        if not links.strip():
            st.error("Pehle kam se kam ek notice link daaliye.")
            return
        pdfs, status = download_pdfs_from_links(links)
        files = []
        for i, (title, data) in enumerate(pdfs, 1):
            name = safe_name(title, f"notice_{i}") + ".pdf"
            files.append((name, data))
        for r in status:
            if r["Result"] == "PDF mil gayi":
                r["Result"] = "OK"
        st.subheader("Status")
        st.dataframe(pd.DataFrame(status), use_container_width=True)
        if files:
            st.success(f"{len(files)} PDF taiyar")
            offer_download(files, "notice_pdfs.zip", "PDF download karo")
        else:
            st.error("Kisi bhi link se PDF nahi mili.")


def ui_resize():
    st.markdown('<div class="hint">PDF, Word (.docx), Excel (.xlsx), JPG, JPEG ya PNG chuniye aur batayiye kitna size chahiye.</div>',
                unsafe_allow_html=True)
    ups = st.file_uploader("File chuniye", type=["pdf", "docx", "xlsx", "jpg", "jpeg", "png"],
                           accept_multiple_files=True, key="up_s")
    c1, c2 = st.columns([2, 1])
    val = c1.number_input("Kitna size chahiye", min_value=1.0, value=100.0, step=10.0, key="val_s")
    unit = c2.selectbox("Unit", list(UNITS), key="unit_s")
    st.caption(f"Bade se chhota karne par compress hoga. Chhote se bada karne par extra data jodkar size badhega (max {human(MAX_PAD)}).")
    if ups and st.button("Size badlo", key="go_s"):
        target = int(val * UNITS[unit])
        outputs, rows = [], []
        for up in ups:
            data = up.read()
            row = {"File": up.name, "Pehle": human(len(data)), "Ab": "", "Result": ""}
            try:
                with st.spinner(f"{up.name} ..."):
                    new, note = resize_file(up.name, data, target)
                outputs.append((up.name, new))
                row["Ab"] = human(len(new))
                row["Result"] = note
            except Exception as e:
                row["Result"] = f"Error: {e}"
            rows.append(row)
        st.subheader("Natija")
        st.dataframe(pd.DataFrame(rows), use_container_width=True)
        offer_download(outputs, "resized_files.zip", "Download karo")


def ui_convert():
    mode = st.radio("Kya badalna hai?", ["Image se Image", "Image se PDF", "PDF se Image"], horizontal=True, key="mode_c")
    img_types = ["jpg", "jpeg", "png", "webp", "bmp", "tif", "tiff"]
    if mode == "Image se Image":
        ups = st.file_uploader("Image chuniye (ek ya zyada)", type=img_types, accept_multiple_files=True, key="up_c1")
        key = st.selectbox("Kis format me badalna hai?", list(IMG_FORMATS), key="fmt_c1")
        if ups and st.button("Convert karo", key="go_c1"):
            outs, rows = [], []
            for up in ups:
                row = {"File": up.name, "Result": ""}
                try:
                    new = convert_image(up.read(), key)
                    outs.append((up.name.rsplit(".", 1)[0] + "." + IMG_FORMATS[key][1], new))
                    row["Result"] = f"OK ({human(len(new))})"
                except Exception as e:
                    row["Result"] = f"Error: {e}"
                rows.append(row)
            st.dataframe(pd.DataFrame(rows), use_container_width=True)
            offer_download(outs, "converted_images.zip", "Download karo")
    elif mode == "Image se PDF":
        ups = st.file_uploader("Images chuniye (jis order me upload karoge, usi order me pages banenge)",
                               type=img_types, accept_multiple_files=True, key="up_c2")
        one = st.checkbox("Sab images ko ek hi PDF me jodo", value=True, key="one_c2")
        if ups and st.button("PDF banao", key="go_c2"):
            try:
                blobs = [(u.name, u.read()) for u in ups]
                if one:
                    outs = [("images.pdf", images_to_pdf([b for _, b in blobs]))]
                else:
                    outs = [(n.rsplit(".", 1)[0] + ".pdf", images_to_pdf([b])) for n, b in blobs]
                st.success(f"{len(outs)} PDF taiyar")
                offer_download(outs, "images_pdf.zip", "PDF download karo")
            except Exception as e:
                st.error(f"Error: {e}")
    else:
        ups = st.file_uploader("PDF chuniye (ek ya zyada)", type=["pdf"], accept_multiple_files=True, key="up_c3")
        c1, c2 = st.columns(2)
        key = c1.selectbox("Image format", ["JPG", "PNG"], key="fmt_c3")
        dpi = c2.selectbox("Quality (dpi)", [100, 150, 200, 300], index=1, key="dpi_c3")
        if ups and st.button("Images banao", key="go_c3"):
            outs, rows = [], []
            for up in ups:
                row = {"File": up.name, "Pages": 0, "Result": ""}
                try:
                    pages = pdf_to_images(up.read(), key, dpi)
                    base = up.name.rsplit(".", 1)[0]
                    for n, ext, b in pages:
                        outs.append((f"{base}_page{n}.{ext}", b))
                    row["Pages"], row["Result"] = len(pages), "OK"
                except Exception as e:
                    row["Result"] = f"Error: {e}"
                rows.append(row)
            st.dataframe(pd.DataFrame(rows), use_container_width=True)
            offer_download(outs, "pdf_pages.zip", "Images download karo")


def ui_excel():
    src = st.radio("Data kahan se aayega?", SRC, horizontal=True, key="src_x")
    if src == SRC[0]:
        ui_links_to_excel()
    else:
        ui_pdf_to_excel()


def ui_word():
    kind = st.radio("Word kaisa chahiye?", ["PDF jaisa layout", "Saaf table"], horizontal=True, key="kind_w")
    if kind == "Saaf table":
        ui_table_to_word()
    else:
        src_w = st.radio("Data kahan se aayega?", SRC, horizontal=True, key="src_w")
        if src_w == SRC[0]:
            ui_links_to_word()
        else:
            ui_pdf_to_word()


def ui_passport():
    st.markdown('<div class="hint">Photo upload kijiye, size aur rang chuniye. Photo katkar sahi size me aayegi aur print ke liye sheet bhi banegi.</div>',
                unsafe_allow_html=True)
    up = st.file_uploader("Photo chuniye", type=["jpg", "jpeg", "png", "webp"], key="up_f")
    if up is None:
        return
    try:
        img = ImageOps.exif_transpose(Image.open(io.BytesIO(up.getvalue())))
        img.load()
    except Exception as e:
        st.error(f"Photo khul nahi payi: {e}")
        return

    c1, c2 = st.columns(2)
    preset = c1.selectbox("Photo ka size", list(PHOTO_SIZES), key="size_f")
    bgname = c2.selectbox("Background rang", list(PHOTO_BG), key="bg_f")
    dims = PHOTO_SIZES[preset]
    if dims is None:
        d1, d2 = st.columns(2)
        dims = (d1.number_input("Chaudai (mm)", 10.0, 200.0, 35.0, 1.0, key="w_f"),
                d2.number_input("Lambai (mm)", 10.0, 200.0, 45.0, 1.0, key="h_f"))
    bg = PHOTO_BG[bgname] or st.color_picker("Rang chuniye", "#FFFFFF", key="bgc_f")

    st.caption("Chehra frame me beech me aaye, isliye zoom aur position se adjust kijiye. Zoom 1 se kam karne par kinare chune hue rang se bhar jate hain.")
    zoom = st.slider("Zoom", 0.5, 3.0, 1.0, 0.05, key="zoom_f")
    s1, s2 = st.columns(2)
    off_x = s1.slider("Left / Right", -100, 100, 0, key="ox_f")
    off_y = s2.slider("Upar / Neeche", -100, 100, 0, key="oy_f")

    photo = make_passport(img, dims[0], dims[1], zoom, off_x, off_y, bg)
    st.image(photo, width=min(300, photo.width // 2),
             caption=f"{dims[0]:g} × {dims[1]:g} mm  ({photo.width} × {photo.height} px, 300 dpi)")

    st.markdown("**Print sheet**")
    p1, p2 = st.columns(2)
    paper = p1.selectbox("Paper", list(PAPERS), key="paper_f")
    guides = p2.checkbox("Katne ke liye line", value=True, key="guides_f")
    _, max_fit = make_sheet(photo, paper, 1, guides)
    allcopies = st.checkbox(f"Jitni aa sake utni (is paper par {max_fit})", value=True, key="all_f")
    copies = None if allcopies else st.number_input("Kitni copies", 1, max(max_fit, 1), min(6, max_fit), key="copies_f")
    sheet, _ = make_sheet(photo, paper, copies, guides)

    limit_on = st.checkbox("Photo ka file size set karna hai (jaise form me 20 se 50 KB)", key="lim_f")
    photo_bytes = _jpg(photo)
    if limit_on:
        kb = st.number_input("Photo ka size (KB)", 5, 5000, 50, 5, key="kb_f")
        try:
            photo_bytes, note = resize_file("photo.jpg", photo_bytes, int(kb * 1024))
            st.caption(f"Photo file: {human(len(photo_bytes))}")
        except Exception as e:
            st.error(f"Error: {e}")

    st.image(sheet, caption=f"{paper} sheet", use_container_width=True)
    d1, d2, d3 = st.columns(3)
    d1.download_button("Photo (JPG)", photo_bytes, "passport_photo.jpg", "image/jpeg", key="dl_f1")
    d2.download_button("Sheet (JPG)", _jpg(sheet), "passport_sheet.jpg", "image/jpeg", key="dl_f2")
    d3.download_button("Sheet (PDF)", _pdf(sheet), "passport_sheet.pdf", "application/pdf", key="dl_f3")


TOOLS = [
    {"id": "excel", "name": "Excel banao", "badge": ".xlsx", "color": "var(--xl)", "fn": ui_excel,
     "desc": "Notice ki roll list ka Excel, links ya PDF se"},
    {"id": "word", "name": "Word banao", "badge": ".docx", "color": "var(--wd)", "fn": ui_word,
     "desc": "PDF jaisa layout ya saaf table, Word me"},
    {"id": "pdf", "name": "Link se PDF", "badge": ".pdf", "color": "var(--pdf)", "fn": ui_links_to_pdf,
     "desc": "Notice ki PDF seedha download karo"},
    {"id": "size", "name": "Size badlo", "badge": "KB<br>MB", "color": "var(--ink)", "fn": ui_resize,
     "desc": "PDF, Word, Excel aur image ka size chhota ya bada karo"},
    {"id": "convert", "name": "Format badlo", "badge": "JPG<br>PNG", "color": "var(--teal)", "fn": ui_convert,
     "desc": "Image aur PDF ko ek dusre me badlo"},
    {"id": "photo", "name": "Passport photo", "badge": "35×<br>45", "color": "var(--amber)", "fn": ui_passport,
     "desc": "Photo sahi size me katkar print sheet banao"},
]
TOOL_BY_ID = {t["id"]: t for t in TOOLS}

if "tool" not in st.session_state:
    st.session_state["tool"] = None


def open_tool(tool_id):
    st.session_state["tool"] = tool_id


def tile_box(key):
    try:
        return st.container(key=key)
    except TypeError:  # purana Streamlit
        return st.container()


def show_home():
    st.markdown(
        """
<div class="hero">
  <div class="files">
    <span class="fs big" style="--c:var(--xl)">.xlsx</span>
    <span class="fs big" style="--c:var(--wd)">.docx</span>
    <span class="fs big" style="--c:var(--pdf)">.pdf</span>
  </div>
  <h1>Notice aur files ke kaam, ek hi jagah</h1>
  <p>Roll list ka Excel ya Word banao, PDF nikalo, file ka size badlo, ya passport photo banao.</p>
</div>
""",
        unsafe_allow_html=True,
    )
    for i in range(0, len(TOOLS), 2):
        cols = st.columns(2)
        for col, t in zip(cols, TOOLS[i:i + 2]):
            with col:
                with tile_box(f"tile_{t['id']}"):
                    st.markdown(
                        f'<div class="tile" style="--c:{t["color"]}"><span class="fs">{t["badge"]}</span>'
                        f'<div><b>{t["name"]}</b><span class="d">{t["desc"]}</span></div></div>',
                        unsafe_allow_html=True,
                    )
                    st.button("Kholo", key=f"open_{t['id']}", on_click=open_tool, args=(t["id"],))


def show_tool(t):
    st.button("← Saare tools", key="back", on_click=open_tool, args=(None,))
    st.markdown(
        f'<div class="thead" style="--c:{t["color"]}"><span class="fs big">{t["badge"]}</span>'
        f'<div><h2>{t["name"]}</h2><p>{t["desc"]}</p></div></div>',
        unsafe_allow_html=True,
    )
    with tile_box("panel"):
        t["fn"]()


current = TOOL_BY_ID.get(st.session_state["tool"])
if current is None:
    show_home()
else:
    show_tool(current)
